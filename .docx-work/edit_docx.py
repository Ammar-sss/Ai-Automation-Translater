import copy
import os

import docx
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

SRC = r"E:\My Projects\program project\Docker-GitHub-CICD-Checklist.docx"
OUT = os.path.join(os.getcwd(), ".docx-work", "Docker-GitHub-CICD-Checklist.updated.docx")

doc = docx.Document(SRC)
body = doc.element.body
blocks = list(body.iterchildren())


def para(idx):
    return Paragraph(blocks[idx], doc)


REF = None
for b in blocks:
    if b.tag == qn("w:p"):
        p = Paragraph(b, doc)
        ppr = b.find(qn("w:pPr"))
        if ppr is not None and ppr.find(qn("w:numPr")) is not None and p.style.name == "List Paragraph":
            REF = p
            break
assert REF is not None
HEAD = para(4).style
LIST = REF.style


def set_text(idx, text):
    p = para(idx)
    p.runs[0].text = text
    for r in p.runs[1:]:
        r.text = ""


def add_p_before(anchor, text, bold_prefix=None):
    p = anchor.insert_paragraph_before("")
    if bold_prefix:
        p.add_run(bold_prefix).bold = True
    p.add_run(text)
    return p


def add_bullet_before(anchor, text):
    p = anchor.insert_paragraph_before("", style=LIST)
    numpr = REF._p.find(qn("w:pPr")).find(qn("w:numPr"))
    p._p.get_or_add_pPr().append(copy.deepcopy(numpr))
    p.add_run(text)
    return p


# ---- Phase ordering: a new phase 7 is inserted, so 7->8 and 8->9 ----
set_heading_text = set_text
set_heading_text(93, "Phase 9: Dependabot and security")
set_heading_text(84, "Phase 8: Run the published image")

anchor = para(84)

# ---- new Phase 7: image scanning ----
anchor.insert_paragraph_before("Phase 7: Scan the image for vulnerabilities", style=HEAD)

add_p_before(anchor, "the base image and the libraries inside it get new CVEs after you build, and nothing tells you unless something looks.", "Why: ")
add_p_before(anchor, "you ship a known-vulnerable web server or JSON library and find out from a scanner, an incident, or never.", "Without it: ")

add_p_before(anchor, "When you create .github/workflows/image-scan.yml, it must:", None)
for t in [
    "Run on every Pull Request, and build the image with push: false and load: true, so a bad image is never published.",
    "Run Trivy (aquasecurity/trivy-action, tag starts with v, e.g. v0.36.0) with severity CRITICAL,HIGH, ignore-unfixed: true and exit-code: '1'.",
    "Set trivyignores: .trivyignore, so every exception is a reviewed line in the repo instead of an invisible one.",
    "Upload the SARIF report to the Security tab with if: always() and continue-on-error: true, because code scanning needs a public repo or GitHub Advanced Security.",
    "Run a second job on a weekly cron that scans the published :latest and only alerts, because new CVEs appear with no code change.",
]:
    add_bullet_before(anchor, t)

add_p_before(anchor, "When you update docker-publish.yml, it must:", None)
for t in [
    "Build with push: false, scan the image, and only then docker push. Pushing first means the vulnerable image is already in Packages.",
    "Scan the commit tag, and keep pushing both the sha tag and latest so rollback still works.",
]:
    add_bullet_before(anchor, t)

add_p_before(anchor, "When you create .trivyignore, it must:", None)
add_bullet_before(anchor, "Give every entry a reason and an exp: date, so it re-surfaces and forces a re-check once a fixed base image exists.")

add_p_before(anchor, "In the branch ruleset: add the scan check as required. Otherwise the gate is decorative.", None)
add_p_before(anchor, "open a Pull Request. A red scan must block the merge button, and the findings must appear under Security.", "Test it: ")
add_p_before(anchor, "read the Fixed Version column and bump that dependency. For Spring-managed libraries that is usually the Spring Boot parent patch version, which moves the web server and Jackson together.", "When the scan is red: ")

# ---- Phase 6: publish only after a scan ----
set_text(82, "Build the image without pushing, scan it, then push the tag and latest only when the scan passed (Phase 7).")

# ---- Phase 4: require the scan check too ----
add_bullet_before(para(56), "Require the image scan check as well once Phase 7 is in place, so a vulnerable image cannot be merged.")

# ---- intro flow line ----
p3 = para(3)
p3.runs[0].text = "The flow: "
for r in p3.runs[1:]:
    r.text = ""
p3.add_run("code on a branch \u2192 Pull Request \u2192 CI builds, tests and scans the image \u2192 merge to main \u2192 CI scans again and publishes the image \u2192 server or PC runs the image.")

# ---- problems table ----
tbl = doc.tables[0]
for row in [
    ("Scan red on a fresh build",
     "CVEs in the base image with no fix released",
     "ignore-unfixed: true; leave the rest to the weekly rescan"),
    ("Trivy action will not resolve in the IDE",
     "Tag written as 0.28.0, but the releases are v0.28.0",
     "Use the v-prefixed tag, or pin a commit SHA"),
    ("PR scan red after a dependency bump",
     "The framework's managed web server and JSON library have known CVEs",
     "Bump the framework parent one patch version"),
    ("Vulnerable image already in Packages",
     "Publish pushed before the scan",
     "Build, scan, then push in the same job"),
    ("Old images fill the disk",
     "Every publish leaves an image behind",
     "docker image prune -f, never -a on a DB host"),
]:
    cells = tbl.add_row().cells
    for c, val in zip(cells, row):
        c.text = val

# ---- final checklist ----
not_done = para(124)
for t in [
    "image-scan.yml committed, and the scan check required in the ruleset",
    "docker-publish builds, scans, then pushes; a red scan publishes nothing",
    "weekly rescan green, or run it once by hand, so new CVEs are caught",
    ".trivyignore entries each have a reason and an expiry date",
]:
    add_bullet_before(not_done, t)

# ---- drop the finished item from "Not done yet" ----
first = para(125)
if first.text.strip().lower().startswith("scan the docker image"):
    first._p.getparent().remove(first._p)

doc.save(OUT)
print("saved:", OUT)

